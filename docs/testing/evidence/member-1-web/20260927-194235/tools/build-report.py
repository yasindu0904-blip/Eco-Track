"""Build the evidence-backed EcoTrack report from the supplied course template.

Run with the Codex document runtime's Python. The original template is read only.
"""

from __future__ import annotations

import csv
import hashlib
import os
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[6]
RUN = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "test-docs" / "6 Template for Test plan.docx"
OUTPUT = RUN / "EcoTrack_Test_Report_DRAFT.docx"
EXPECTED_SHA = "471dcb5496001b91a0bb0cd59cb851439f16e467777dcc95dbcde2bc0ef861de"

if hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() != EXPECTED_SHA:
    raise RuntimeError("The course template changed; distill it again before building")

doc = Document(TEMPLATE)
doc.core_properties.title = "EcoTrack Test Report"
doc.core_properties.subject = "Web functional testing and combined system evidence"
doc.core_properties.author = "Sahanya"
doc.core_properties.last_modified_by = "Sahanya"


def drop(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def paragraphs_matching(text):
    return [paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == text]


def anchor(text):
    matches = paragraphs_matching(text)
    if not matches:
        raise RuntimeError(f"Template heading missing: {text}")
    return matches[0]


def add_before(next_heading, content, *, style="Body Text"):
    return anchor(next_heading).insert_paragraph_before(content, style=style)


def fill_technique(table, objective, method, oracle, tools, criterion, note):
    for row, value in zip(table.rows, [objective, method, oracle, tools, criterion, note]):
        cell = row.cells[1]
        cell.text = value
        for paragraph in cell.paragraphs:
            paragraph.style = doc.styles["Body Text"]
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)


for section in doc.sections:
    for paragraph in section.header.paragraphs:
        if "<Company Name>" in paragraph.text:
            paragraph.text = paragraph.text.replace("<Company Name>", "EcoTrack")
    for paragraph in section.footer.paragraphs:
        if "<Company Name>" in paragraph.text:
            paragraph.text = paragraph.text.replace("<Company Name>", "EcoTrack")

for paragraph in doc.paragraphs:
    if paragraph.text.strip() == "<Project Name>":
        paragraph.text = "EcoTrack"
    elif "<Iteration/ Master> Test Plan" in paragraph.text:
        paragraph.text = "Web and System Test Report"
    elif paragraph.text.strip() == "Version <1.0>":
        paragraph.text = "Version 0.1 — 27 September 2026"

revision = doc.tables[0]
for cell, value in zip(revision.rows[1].cells, ["27/Sep/26", "0.1", "Evidence-backed draft; role workflows pending", "Sahanya"]):
    cell.text = value
for row in revision.rows[2:]:
    for cell in row.cells:
        cell.text = ""

toc = [paragraph for paragraph in doc.paragraphs if paragraph.style.name.startswith("toc ")]
toc_lines = [
    "1. Evaluation Mission and Test Motivation",
    "2. Target Test Items",
    "3. Test Approach",
    "4. Actual Results and Evidence",
    "5. Deliverables",
    "6. Risks, Dependencies, Assumptions, and Constraints",
    "7. References",
]
for paragraph in toc:
    drop(paragraph)
toc_cursor = anchor("Table of Contents")
for line in toc_lines:
    element = OxmlElement("w:p")
    toc_cursor._p.addnext(element)
    toc_cursor = Paragraph(element, toc_cursor._parent)
    toc_cursor.style = doc.styles["Body Text"]
    toc_cursor.add_run(line)

for paragraph in list(doc.paragraphs):
    if paragraph.style.name == "InfoBlue":
        drop(paragraph)
    elif paragraph.text.startswith("The listing below identifies those test items"):
        drop(paragraph)
    elif paragraph.text.startswith("Indicate the tool references"):
        drop(paragraph)
    elif paragraph.text.startswith("Refer any data/"):
        drop(paragraph)
    elif paragraph.text.startswith("For different algorithms/"):
        drop(paragraph)
    elif paragraph.text.startswith("For tools you can"):
        drop(paragraph)
    elif paragraph.text.startswith("For similar work"):
        drop(paragraph)
    elif paragraph.text.startswith("You may include white paper"):
        drop(paragraph)
    elif paragraph.text.strip() == "6.      References":
        drop(paragraph)
    elif paragraph.style.name == "List Paragraph" and not paragraph.text.strip():
        drop(paragraph)

add_before("Target Test Items", "EcoTrack is a multi-tenant platform for citizen incident reporting and organization-led cleanup. This evaluation examines whether the hosted web experience supports the principal workflows and protects private records. The report records actual observations and outstanding risks for the 27 September 2026 submission.")
add_before("Target Test Items", "The team divides execution by surface: Sahanya owns web functional/UI testing and report assembly; Member 2 owns API, database, security, performance and deployment diagnosis; Member 3 owns Android device, resource and notification evidence. Usability participant testing is outside these assignments.")

add_before("Test Approach", "The test items are the hosted EcoTrack web app, hosted API, Supabase Auth/Storage, the disposable application test database, official Kesbewa geography, and representative Windows browser configurations. Hosted web: https://eco-track-ctd.pages.dev. Hosted API: https://p01--eco-track--xnjn7t9fk69n.code.run/api/v1.")
add_before("Test Approach", "Local source revision: 4181db9c2dd455569aa872ce0c73541c8cd632ce. The deployed revision has not been verified; local source is used for expected behavior and local checks, not as proof of the hosted build.")

add_before("Testing Techniques and Types", "W01–W17 cover authentication, profile completion, incident creation and persistence, map discovery, organization application/review, event lifecycle, notifications, keyboard/narrow-width use, browser comparison, offline recovery and sign-out. P0 flows are prioritized. Every case starts NOT_RUN and moves to PASS, FAIL or BLOCKED only after an attempted execution is recorded.")
add_before("Testing Techniques and Types", "Oracles are current source validation, the team guide's expected outcomes, controlled fixture IDs, visible UI states, sanitized request/response status and content, and persisted records when checked by Member 2. Empty lists and HTTP 200 responses alone do not prove filtering or recipient correctness.")

fill_technique(doc.tables[1], "Verify persistence and tenant integrity of application data.", "Member 2 runs isolated API/SQL checks against a verified test DB; Member 1 references results and UI persistence only.", "Returned values, known fixture IDs, database rows and tenant boundaries agree.", "Member 2 tools and versions: pending handover.", "Key integrity and cross-tenant assertions observed; setup excluded from counts.", "Do not run destructive integration suites against the shared hosted test DB.")
fill_technique(doc.tables[2], "Exercise citizen and organization business flows.", "Run W01–W13 and W17 through hosted web with separate role contexts and synthetic records.", "Visible final state, API result and persisted details match each case expectation.", "Playwright Core 1.55.1, Edge 154, Chrome 153; local Vitest 4.1.11.", "Each full case has timestamped evidence and a justified PASS/FAIL/BLOCKED verdict.", "UI setup actions and SQL/API fixture setup are labeled separately.")
fill_technique(doc.tables[3], "Check navigation, validation, focus and narrow viewport behavior.", "Use semantic controls, keyboard Tab/Shift+Tab/Enter/Escape, 1440x1000 and 390x844 views, and screenshots of loaded states.", "Fields, errors, focus, modal controls and layout are understandable and operable.", "Playwright screenshots; Edge/Chrome developer inspection where available.", "No clipped core controls or page horizontal overflow; issues documented with exact scope.", "Basic UI checks are not a SUS participant study or blanket accessibility claim.")
fill_technique(doc.tables[4], "Measure individual response times for selected operations.", "Member 2 captures timed API requests after an agreed quiet window.", "Observed timestamps and response status/content are recorded per request.", "Member 2 measurement tools: pending handover.", "Actual timings and environmental conditions are reported without inventing an SLA.", "Smoke request durations alone are not load-test evidence.")
fill_technique(doc.tables[5], "Evaluate response under concurrent load.", "Member 2 runs the agreed controlled load plan against the test target.", "Error rate, latency distribution and server metrics correspond to the same window.", "Member 2 tools/results: pending handover.", "Only executed load results and known limits appear in the report.", "No load run overlaps normal team interaction or shared service changes.")
fill_technique(doc.tables[6], "Verify authenticated role and record boundaries.", "Member 1 checks browser sign-out and cross-account visibility; Member 2 checks API authorization and tenant IDs.", "Unauthorized access rejected and private records absent outside owning context.", "Playwright and hosted API observations; Member 2 security tools pending.", "Positive and negative role assertions have separate evidence.", "Route hiding alone does not establish backend authorization.")
fill_technique(doc.tables[7], "Observe controlled failure and recovery.", "W16 toggles browser offline mode for a read and controlled save, then retries after reconnection.", "Clear error/recovery without duplicate writes or loss of confirmed data.", "Edge/Chrome DevTools or Playwright network emulation, if executed.", "Failure and recovery are captured as separate states.", "Do not claim failover of backend infrastructure from a browser offline test.")
fill_technique(doc.tables[8], "Compare representative browser and device configurations.", "W15 compares Edge and Chrome; Member 3 checks Android device/build behavior.", "Core tasks render and behave consistently in observed configurations.", "Edge 154, Chrome 153; Android details pending Member 3.", "Report exact versions and observed combinations only.", "Edge and Chrome share Chromium; they are not independent rendering engines.")

actual = anchor("Deliverables")
actual.insert_paragraph_before("Actual Results and Evidence", style="Heading 1")
actual.insert_paragraph_before("Member 1 hosted web", style="Heading 2")
with (RUN / "cases.csv").open(newline="", encoding="utf-8") as handle:
    cases = list(csv.DictReader(handle))
counts = {status: sum(item["status"] == status for item in cases) for status in ["PASS", "FAIL", "BLOCKED", "NOT_RUN"]}
actual.insert_paragraph_before(
    f"At this draft checkpoint: {len(cases)} planned web cases; {counts['PASS']} PASS, {counts['FAIL']} FAIL, {counts['BLOCKED']} BLOCKED, {counts['NOT_RUN']} NOT_RUN. Partial observations do not count as complete case passes.",
    style="Body Text",
)
actual.insert_paragraph_before("Fresh hosted observations: Edge loaded the signed-out login and blocked a malformed email. Tab and Shift+Tab followed the email-submit order; keyboard Enter on malformed input triggered native validation without an auth API request. Keyboard focus was visible and the 390 px login page had no horizontal document overflow. Chrome rendered the signed-out login. Unauthenticated GET /api/v1/auth/me returned HTTP 401. Magic-link callback and role workflows remain pending in this draft.", style="Body Text")
actual.insert_paragraph_before("Detailed case steps, each attempt, UTC timestamps, actual statuses and sanitized paths are in the Member 1 run's cases.csv, results.csv and evidence-index.csv. These are the authoritative execution records.", style="Body Text")
actual.insert_paragraph_before("Local automated checks", style="Heading 2")
actual.insert_paragraph_before("Vitest passed 48 tests in 16 files. TypeScript/Vite build passed (187 modules). ESLint passed on a larger-heap retry. The initial npm lint attempt exhausted Node memory and the initial Vitest startup hit spawn EPERM; both are retained as tool failures. Local checks do not establish hosted acceptance.", style="Body Text")
actual.insert_paragraph_before("Prior hosted smoke baseline", style="Heading 2")
actual.insert_paragraph_before("The separate 27 September smoke package records an initial hosted incident POST HTTP 500 at 12:35:06.530 UTC and a later successful retest without an application fix. D01 remains open with an unconfirmed cause. The lifecycle completion HTTP 409 came from a synthetic fixture's timestamp precision; it passed after fixture correction. A browser locator expected Unread text while All was selected; visual review found the correct empty inbox. These observations are cited under their original run and not counted as new Member 1 cases.", style="Body Text")
actual.insert_paragraph_before("Pending cross-member handovers", style="Heading 2")
actual.insert_paragraph_before("Member 2 API/database/security/performance and Member 3 Android/device/notification result packages are not yet present in the repository. Their cases remain pending in this report. Real phone receipt, email delivery, registration/approval, populated spatial filtering and full browser event lifecycle must not be inferred from the prior smoke package.", style="Body Text")

caption = actual.insert_paragraph_before("Figure 1. Malformed email blocked on hosted login. Evidence: member-1-web/20260927-194235/screenshots/W01_attempt-01_invalid-email.png (27 September 2026, 14:26:50 UTC).", style="Body Text")
caption.paragraph_format.keep_with_next = True
figure = actual.insert_paragraph_before(style="Body Text")
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.add_run().add_picture(str(RUN / "screenshots" / "W01_attempt-01_invalid-email.png"), width=Inches(4.8))
caption = actual.insert_paragraph_before("Figure 2. Hosted login at 390x844. Evidence: member-1-web/20260927-194235/screenshots/W14_attempt-01_narrow.png (27 September 2026, 14:26:50 UTC).", style="Body Text")
caption.paragraph_format.keep_with_next = True
figure = actual.insert_paragraph_before(style="Body Text")
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.add_run().add_picture(str(RUN / "screenshots" / "W14_attempt-01_narrow.png"), height=Inches(3.6))

add_before("Test Evaluation Summaries", "The Member 1 evidence folder contains environment metadata, planned cases, per-attempt results, fixture manifest, defects, screenshot index, sanitized logs/requests and reproducible tool snapshots. The existing smoke run is retained at its original path. A combined final report requires the other members' verified handovers.")
add_before("Reporting on Test Coverage", "This report summarizes unique W-case verdicts, links attempts and retests separately, and names unresolved defects. It excludes setup, cleanup and repeated attempts from product-test counts.")
add_before("Risks, Dependencies, Assumptions, and Constraints", "The authoritative per-attempt record is results.csv. Source review, local Vitest/lint/build, hosted browser actions, prior smoke, and human observations are presented as distinct evidence sources. No pass rate is reported until the planned suite is genuinely executed.")

risks = [
    ("Test database identity", "Verify hosted API points at eco-track-test before writes.", "Pause write cases; continue read-only UI and local checks."),
    ("Role access", "Use volunteer, org-admin and superadmin reviewer test accounts.", "Mark dependent cases BLOCKED and retain partial evidence."),
    ("Incident HTTP 500", "Keep D01 open; request server logs and distinct retest.", "Report intermittent failure and limit acceptance claim."),
]
for row, values in zip(doc.tables[9].rows[1:], risks):
    for cell, value in zip(row.cells, values):
        cell.text = value

refs = [
    "1. EcoTrack Member 1 web testing guide, repository, accessed 27 September 2026.",
    "2. EcoTrack 2026-09-27 hosted smoke README, RESULTS.csv, defects and harness notes, accessed 27 September 2026.",
    "3. EcoTrack local source at 4181db9c2dd455569aa872ce0c73541c8cd632ce, accessed 27 September 2026.",
    "4. Course-supplied 6 Template for Test plan.docx, accessed 27 September 2026.",
    "5. Playwright Page API: https://playwright.dev/docs/api/class-page, accessed 27 September 2026.",
    "6. Vitest guide: https://vitest.dev/guide/, accessed 27 September 2026.",
    "7. Vite guide: https://vite.dev/guide/, accessed 27 September 2026.",
    "8. ESLint docs: https://eslint.org/docs/latest/, accessed 27 September 2026.",
]
for item in refs:
    doc.add_paragraph(item, style="Body Text")
anchor("References").paragraph_format.page_break_before = True

for paragraph in list(doc.paragraphs):
    if paragraph.style.name == "InfoBlue":
        raise RuntimeError("Template guidance survived")

for node in doc.part.element.xpath(".//w:t"):
    if node.text:
        node.text = node.text.replace("<Iteration/ Master> Test Plan", "Test Report")
for section in doc.sections:
    for node in section.header.part.element.xpath(".//w:t"):
        if node.text:
            node.text = node.text.replace("<Company Name>", "EcoTrack")
    for node in section.footer.part.element.xpath(".//w:t"):
        if node.text:
            node.text = node.text.replace("<Company Name>", "EcoTrack")

doc.save(OUTPUT)
temporary = OUTPUT.with_suffix(".tmp.docx")
with ZipFile(OUTPUT, "r") as source, ZipFile(temporary, "w") as destination:
    for item in source.infolist():
        data = source.read(item.filename)
        if item.filename == "docProps/app.xml":
            data = data.replace(b"&lt;Company Name&gt;", b"EcoTrack")
            data = data.replace(b"&lt;Iteration/ Master&gt; Test Plan", b"EcoTrack Test Report")
        destination.writestr(item, data)
os.replace(temporary, OUTPUT)
if hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() != EXPECTED_SHA:
    raise RuntimeError("Template changed during build")
print(OUTPUT)
